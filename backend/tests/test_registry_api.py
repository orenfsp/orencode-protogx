from datetime import date, timedelta
from io import BytesIO

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import (
    AthleteProfile,
    Guardian,
    PersonalDataConsent,
    User,
    UserRole,
)
from apps.catalogs.models import Discipline, Organization
from apps.competitions.models import AgeGroup, AuditEvent
from apps.registry.models import (
    DisciplinePeriod,
    MunicipalityPeriod,
    OrganizationPeriod,
    ParticipationPeriod,
)


class RegistryApiTests(TestCase):
    def setUp(self):
        self.federation = User.objects.create_user(email="f@example.test", full_name="Federation")
        UserRole.objects.create(user=self.federation, role="federation")
        self.athlete = User.objects.create_user(email="a@example.test", full_name="Fictional Athlete", date_of_birth=date(2000, 1, 1))
        UserRole.objects.create(user=self.athlete, role="athlete")
        self.first = Organization.objects.create(code="first-api", name="Первая")
        self.second = Organization.objects.create(code="second-api", name="Вторая")
        AthleteProfile.objects.create(user=self.athlete, organization=self.second)
        a = Discipline.objects.create(code="api-a", name="Первая дисциплина")
        b = Discipline.objects.create(code="api-b", name="Вторая дисциплина")
        AgeGroup.objects.create(name="Взрослые тест", min_age=18, max_age=40)
        ParticipationPeriod.objects.create(athlete=self.athlete, starts_on=date(2024, 1, 1))
        OrganizationPeriod.objects.create(athlete=self.athlete, organization=self.first, starts_on=date(2024, 1, 1), ends_on=date(2024, 12, 31))
        OrganizationPeriod.objects.create(athlete=self.athlete, organization=self.second, starts_on=date(2025, 1, 1))
        MunicipalityPeriod.objects.create(athlete=self.athlete, municipality="Оренбург", starts_on=date(2024, 1, 1))
        for item in (a, b):
            DisciplinePeriod.objects.create(athlete=self.athlete, discipline=item, starts_on=date(2024, 1, 1))
        self.client = APIClient()
        self.client.force_authenticate(self.federation)
        session = self.client.session
        session["active_role"] = "federation"
        session.save()

    def test_unique_and_historical_breakdowns_and_xlsx(self):
        url = "/api/v1/analytics/participation/?start_year=2024&end_year=2025"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Cache-Control"], "no-store")
        self.assertEqual([row["unique_athletes"] for row in response.data], [1, 1])
        self.assertEqual(len(response.data[0]["breakdowns"]["disciplines"]), 2)
        self.assertEqual(response.data[0]["breakdowns"]["organizations"][0]["name"], "Первая")
        self.assertEqual(response.data[1]["breakdowns"]["organizations"][0]["name"], "Вторая")
        export = self.client.get(url.replace("participation/?", "participation/export/?"))
        self.assertEqual(export["X-Report-Fingerprint"], response["X-Report-Fingerprint"])
        book = load_workbook(BytesIO(export.content), data_only=True)
        self.assertIn(response["X-Report-Fingerprint"], book["Методика"]["A1"].value)
        self.assertEqual(book["Итоги"]["B4"].value, response.data[0]["unique_athletes"])
        self.assertEqual(book["Организации"]["B2"].value, "Первая")
        self.assertEqual(book["Организации"]["B3"].value, "Вторая")
        filtered = self.client.get(url + f"&organization={self.first.pk}")
        self.assertEqual([row["unique_athletes"] for row in filtered.data], [1, 0])
        filtered_file = self.client.get(url.replace("participation/?", "participation/export/?") + f"&organization={self.first.pk}")
        filtered_book = load_workbook(BytesIO(filtered_file.content), data_only=True)
        self.assertEqual([filtered_book["Итоги"][f"B{row}"].value for row in (4, 5)], [1, 0])
        self.assertEqual(filtered_file["X-Report-Fingerprint"], filtered["X-Report-Fingerprint"])

    def test_federation_registry_and_role_denial(self):
        for path in ("athletes", "coaches", "judges", "organizations"):
            self.assertEqual(self.client.get(f"/api/v1/{path}/").status_code, 200)
        self.assertEqual(self.client.get("/api/v1/athletes/?search=Fictional").data["count"], 1)
        self.client.force_authenticate(self.athlete)
        session = self.client.session
        session["active_role"] = "athlete"
        session.save()
        self.assertEqual(self.client.get("/api/v1/athletes/").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/analytics/participation/export/?start_year=2024&end_year=2025").status_code, 403)

    def test_filters_must_overlap_at_the_same_time(self):
        before = Discipline.objects.create(code="before-transfer", name="До перевода")
        DisciplinePeriod.objects.create(athlete=self.athlete, discipline=before, starts_on=date(2024, 1, 1), ends_on=date(2024, 6, 30))
        OrganizationPeriod.objects.filter(athlete=self.athlete, organization=self.first).update(ends_on=date(2024, 6, 30))
        OrganizationPeriod.objects.filter(athlete=self.athlete, organization=self.second).update(starts_on=date(2024, 7, 1))
        result = self.client.get(f"/api/v1/analytics/participation/?start_year=2024&end_year=2025&organization={self.second.pk}&discipline={before.pk}")
        self.assertEqual([row["unique_athletes"] for row in result.data], [0, 0])

    def test_invalid_query_and_organization_code_are_rejected(self):
        self.assertEqual(self.client.get("/api/v1/analytics/participation/?start_year=2025&end_year=2025").status_code, 400)
        response = self.client.patch(f"/api/v1/organizations/{self.first.pk}/", {"code": "changed"}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_correction_is_audited_and_stale_or_unauthorized_rejected(self):
        period = ParticipationPeriod.objects.get(athlete=self.athlete)
        path = f"/api/v1/athletes/{self.athlete.pk}/history/participation/{period.pk}/correct/"
        body = {"expected": {"starts_on": "2024-01-01", "ends_on": None}, "corrected": {"starts_on": "2025-01-01", "ends_on": None}, "reason": "Уточнена дата"}
        self.assertEqual(self.client.post(path, body, format="json").status_code, 200)
        event = AuditEvent.objects.get(action="registry.period_corrected")
        self.assertEqual(event.before, body["expected"])
        self.assertEqual(event.after["values"], body["corrected"])
        self.assertEqual(event.after["reason"], "Уточнена дата")
        audit_path = f"/api/v1/athletes/{self.athlete.pk}/history/participation/{period.pk}/audit/"
        self.assertEqual(self.client.get(audit_path).data[0]["after"]["reason"], "Уточнена дата")
        self.assertEqual(self.client.post(path, body, format="json").status_code, 409)
        report = self.client.get("/api/v1/analytics/participation/?start_year=2024&end_year=2025")
        self.assertEqual([row["unique_athletes"] for row in report.data], [0, 1])
        self.client.force_authenticate(self.athlete)
        session = self.client.session
        session["active_role"] = "athlete"
        session.save()
        for role in ("athlete", "coach", "judge"):
            self.assertEqual(self.client.post(f"/api/v1/registry/profiles/{role}/", body, format="json").status_code, 403)
        self.assertEqual(self.client.get(audit_path).status_code, 403)
        self.assertEqual(AuditEvent.objects.filter(action="registry.period_corrected").count(), 1)

    def test_overlapping_correction_is_rejected_without_audit(self):
        existing = ParticipationPeriod.objects.get(athlete=self.athlete)
        existing.ends_on = date(2024, 12, 31)
        existing.save()
        other = ParticipationPeriod.objects.create(athlete=self.athlete, starts_on=date(2025, 1, 1))
        path = f"/api/v1/athletes/{self.athlete.pk}/history/participation/{other.pk}/correct/"
        body = {"expected": {"starts_on": "2025-01-01", "ends_on": None}, "corrected": {"starts_on": "2024-12-31", "ends_on": None}, "reason": "Correction"}
        self.assertEqual(self.client.post(path, body, format="json").status_code, 400)
        self.assertFalse(AuditEvent.objects.filter(action="registry.period_corrected").exists())

    def test_profile_mutations_are_federation_only_for_each_role(self):
        people = {"athletes": self.athlete}
        for kind, role in (("coaches", "coach"), ("judges", "judge")):
            user = User.objects.create_user(email=f"{role}@example.test", full_name=f"Original {role}")
            UserRole.objects.create(user=user, role=role)
            people[kind] = user
        for kind, user in people.items():
            path = f"/api/v1/{kind}/{user.pk}/"
            response = self.client.patch(path, {"full_name": f"Corrected {kind}"}, format="json")
            self.assertEqual(response.status_code, 200)
        self.assertEqual(AuditEvent.objects.filter(action="registry.person_updated").count(), 3)
        self.client.force_authenticate(self.athlete)
        session = self.client.session
        session["active_role"] = "athlete"
        session.save()
        for kind, user in people.items():
            self.assertEqual(self.client.patch(f"/api/v1/{kind}/{user.pk}/", {"full_name": "Unauthorized"}, format="json").status_code, 403)
        self.assertEqual(self.client.patch(f"/api/v1/organizations/{self.first.pk}/", {"name": "Unauthorized"}, format="json").status_code, 403)
        self.assertEqual(self.client.post("/api/v1/organizations/", {"code": "bad", "name": "Unauthorized"}, format="json").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/registry/options/").status_code, 403)

    def test_registry_filters_and_reference_options(self):
        response = self.client.get(f"/api/v1/athletes/?organization={self.second.pk}&active=true")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(self.client.get(f"/api/v1/athletes/?organization={self.first.pk}").data["count"], 0)
        self.assertEqual(self.client.get("/api/v1/athletes/?active=invalid").status_code, 400)
        options = self.client.get("/api/v1/registry/options/")
        self.assertEqual(options.status_code, 200)
        self.assertIn("Оренбург", options.data["municipalities"])
        self.assertIn({"id": self.second.pk, "name": "Вторая"}, options.data["organizations"])

    def test_organization_filters_validate_and_combine(self):
        self.first.municipality = "Оренбург"
        self.first.organization_type = "Клуб"
        self.first.save(update_fields=["municipality", "organization_type"])
        self.second.municipality = "Орск"
        self.second.is_active = False
        self.second.save(update_fields=["municipality", "is_active"])
        response = self.client.get("/api/v1/organizations/", {"active": "true", "municipality": "Оренбург", "organization_type": "Клуб"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([item["id"] for item in response.data["results"]], [self.first.pk])
        self.assertEqual(self.client.get("/api/v1/organizations/", {"active": "false"}).data["count"], 1)
        self.assertEqual(self.client.get("/api/v1/organizations/", {"active": "invalid"}).status_code, 400)

    def test_session_mutation_requires_csrf(self):
        client = APIClient(enforce_csrf_checks=True)
        client.force_login(self.federation)
        session = client.session
        session["active_role"] = "federation"
        session.save()
        response = client.post("/api/v1/organizations/", {"code": "csrf-test", "name": "Тест"}, format="json")
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Organization.objects.filter(code="csrf-test").exists())

    def test_profile_creation_requires_existing_role_and_consent(self):
        from apps.registry.models import StaffProfile

        coach = User.objects.create_user(email="coach-new@example.test", full_name="Coach", date_of_birth=date(1980, 1, 1))
        UserRole.objects.create(user=coach, role="coach")
        path = "/api/v1/registry/profiles/coach/"
        body = {"user": coach.pk, "organization": self.first.pk}
        self.assertEqual(self.client.post(path, body, format="json").status_code, 400)
        PersonalDataConsent.objects.create(user=coach, text_version="test")
        self.assertEqual(self.client.post(path, body, format="json").status_code, 201)
        self.assertEqual(StaffProfile.objects.get(user=coach, role="coach").organization_id, self.first.pk)
        discipline = Discipline.objects.get(code="api-a")
        self.first.municipality = "Орск"
        self.first.save(update_fields=["municipality"])
        staff_path = f"/api/v1/registry/profiles/coach/{coach.pk}/staff/"
        self.assertEqual(self.client.patch(staff_path, {"disciplines": [discipline.pk]}, format="json").status_code, 200)
        self.assertEqual(self.client.get(f"/api/v1/coaches/?discipline={discipline.pk}&municipality=Орск").data["count"], 1)
        self.assertTrue(AuditEvent.objects.filter(action="registry.staff_profile_updated").exists())
        self.assertEqual(self.client.post(path, body, format="json").status_code, 400)
        self.assertEqual(self.client.post("/api/v1/registry/profiles/judge/", body, format="json").status_code, 400)
        self.assertFalse(coach.role_assignments.filter(role="judge").exists())
        judge = User.objects.create_user(email="judge-create@example.test", full_name="Judge", date_of_birth=date(1980, 1, 1))
        UserRole.objects.create(user=judge, role="judge")
        PersonalDataConsent.objects.create(user=judge, text_version="test")
        self.assertEqual(self.client.post("/api/v1/registry/profiles/judge/", {"user": judge.pk}, format="json").status_code, 201)
        self.assertTrue(StaffProfile.objects.filter(user=judge, role="judge").exists())
        self.client.force_authenticate(self.athlete)
        session = self.client.session
        session["active_role"] = "athlete"
        session.save()
        self.assertEqual(self.client.post(path, body, format="json").status_code, 403)

    def test_minor_profile_requires_verified_guardian_on_current_consent(self):
        junior = User.objects.create_user(email="junior-new@example.test", full_name="Junior", date_of_birth=date(2015, 1, 1))
        UserRole.objects.create(user=junior, role="athlete")
        guardian = Guardian.objects.create(minor=junior, full_name="Fictional guardian", relationship="parent")
        PersonalDataConsent.objects.create(user=junior, guardian=guardian, text_version="test")
        path = "/api/v1/registry/profiles/athlete/"
        self.assertEqual(self.client.post(path, {"user": junior.pk}, format="json").status_code, 400)
        guardian.confirmation_status = "verified"
        guardian.save(update_fields=["confirmation_status"])
        self.assertEqual(self.client.post(path, {"user": junior.pk}, format="json").status_code, 201)
        self.assertTrue(AthleteProfile.objects.filter(user=junior).exists())

    def test_context_and_staff_update_are_restricted(self):
        from apps.registry.models import StaffProfile

        judge = User.objects.create_user(email="judge-new@example.test", full_name="Judge")
        UserRole.objects.create(user=judge, role="judge")
        StaffProfile.objects.create(user=judge, role="judge")
        staff = f"/api/v1/registry/profiles/judge/{judge.pk}/staff/"
        self.assertEqual(self.client.patch(staff, {"organization": self.first.pk}, format="json").status_code, 200)
        self.assertEqual(self.client.get(f"/api/v1/registry/profiles/judge/{judge.pk}/context/").status_code, 200)
        self.client.force_authenticate(self.athlete)
        session = self.client.session
        session["active_role"] = "athlete"
        session.save()
        self.assertEqual(self.client.get(f"/api/v1/registry/profiles/judge/{judge.pk}/context/").status_code, 403)
        self.assertEqual(self.client.patch(staff, {"organization": self.second.pk}, format="json").status_code, 403)

    def test_judge_context_marks_current_category_and_expiry(self):
        from apps.qualifications.models import (
            JudgeCategoryApplication,
            JudgeCategoryHistory,
        )

        judge = User.objects.create_user(email="judge-category@example.test", full_name="Judge")
        UserRole.objects.create(user=judge, role="judge")
        today = timezone.localdate()
        for category, assigned, valid_until in (
            ("third", today - timedelta(days=800), today - timedelta(days=1)),
            ("second", today - timedelta(days=100), today + timedelta(days=400)),
        ):
            application = JudgeCategoryApplication.objects.create(judge=judge, request_type="assign", requested_category=category, status="approved")
            JudgeCategoryHistory.objects.create(judge=judge, category=category, application=application, order_number="Fictional", assigned_on=assigned, valid_until=valid_until, recorded_by=self.federation)
        response = self.client.get(f"/api/v1/registry/profiles/judge/{judge.pk}/context/")
        self.assertEqual(response.status_code, 200)
        current = [item for item in response.data["categories"] if item["current"]]
        self.assertEqual(len(current), 1)
        self.assertEqual(current[0]["validity"], "active")
        self.assertEqual(current[0]["valid_until"], today + timedelta(days=400))
        self.assertTrue(any(item["validity"] == "expired" and not item["current"] for item in response.data["categories"]))
