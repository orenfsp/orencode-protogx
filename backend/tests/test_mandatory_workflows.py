from datetime import date
from io import BytesIO

from django.test import TestCase
from docx import Document
from openpyxl import load_workbook

from apps.accounts.models import Guardian, PersonalDataConsent, has_current_consent
from apps.competitions.models import AuditEvent, CoachAthleteLink
from tests import test_competitions as competition_tests


class MandatoryWorkflowTests(TestCase):
    setUp = competition_tests.CompetitionWorkflowTests.setUp
    user = competition_tests.CompetitionWorkflowTests.user
    as_role = competition_tests.CompetitionWorkflowTests.as_role
    competition = competition_tests.CompetitionWorkflowTests.competition
    apply = competition_tests.CompetitionWorkflowTests.apply
    approve = competition_tests.CompetitionWorkflowTests.approve

    def test_coach_directory_minimal_consent_minor_and_role_scope(self):
        self.athlete.full_name = "Вымышленный Спортсмен"
        self.athlete.save()
        self.second.full_name = "Вымышленный Несовершеннолетний"
        self.second.date_of_birth = date(2015, 1, 1)
        self.second.save()
        self.as_role(self.coach, "coach")
        path = "/api/v1/coach-athlete-directory/"
        self.assertEqual(self.client.get(path + "?search=В").status_code, 400)
        response = self.client.get(path + "?search=Вымышленный")
        self.assertEqual(response.data, [{"id": self.athlete.pk, "full_name": self.athlete.full_name}])
        data = {"athlete": self.second.pk, "discipline": self.algorithmic.pk, "search": "Вымышленный"}
        self.assertEqual(self.client.post(path, data).status_code, 403)
        data["athlete"] = self.athlete.pk
        self.assertEqual(self.client.post(path, data).status_code, 202)
        self.assertEqual(self.client.post(path, data).status_code, 202)
        link = CoachAthleteLink.objects.get(coach=self.coach, athlete=self.athlete)
        self.assertEqual(link.status, "pending")
        consent = PersonalDataConsent.objects.get(user=self.athlete)
        from django.utils import timezone
        consent.revoked_at = timezone.now()
        consent.save()
        self.assertEqual(self.client.get(path + "?search=Вымышленный").data, [])
        self.assertEqual(self.client.post(path, data).status_code, 403)
        self.as_role(self.athlete, "athlete")
        self.assertEqual(self.client.get(path + "?search=Вымышленный").status_code, 403)

    def test_link_history_both_owners_only_and_records_reasons(self):
        self.as_role(self.coach, "coach")
        self.client.post("/api/v1/coach-links/", {"athlete_email": self.athlete.email, "discipline": self.algorithmic.pk})
        link = CoachAthleteLink.objects.get(coach=self.coach, athlete=self.athlete)
        self.as_role(self.athlete, "athlete")
        self.client.post(f"/api/v1/coach-links/{link.pk}/decision/", {"status": "confirmed"})
        self.client.post(f"/api/v1/coach-links/{link.pk}/status/", {"status": "suspended", "reason": "Учебный перерыв"})
        path = f"/api/v1/coach-links/{link.pk}/history/"
        rows = self.client.get(path).data
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[-1]["before"]["status"], "confirmed")
        self.assertEqual(rows[-1]["after"]["reason"], "Учебный перерыв")
        self.as_role(self.coach, "coach")
        self.assertEqual(len(self.client.get(path).data), 3)
        self.as_role(self.second, "athlete")
        self.assertEqual(self.client.get(path).status_code, 404)

    def test_personal_protocol_hides_other_rows_and_has_version(self):
        competition = self.competition()
        for athlete in (self.athlete, self.second):
            application = self.apply(competition, athlete)
            self.approve(application.data["id"])
        self.client.post(f"/api/v1/competitions/{competition}/protocol-publish/")
        self.as_role(self.athlete, "athlete")
        path = f"/api/v1/history/{competition}/protocols/1/?file_format=docx"
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200, response.content[:50])
        doc = Document(BytesIO(response.content))
        contents = "\n".join(p.text for p in doc.paragraphs) + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn(self.athlete.full_name, contents)
        self.assertNotIn(self.second.full_name, contents)
        self.assertIn("Личная выписка", contents)
        self.as_role(self.coach, "coach")
        self.assertEqual(self.client.get(path).status_code, 403)

    def test_http_errors_are_russian_with_stable_status_and_field_names(self):
        self.as_role(self.athlete, "athlete")
        response = self.client.post("/api/v1/competitions/", {}, format="json")
        self.assertEqual(response.status_code, 403)
        self.assertIn("роль", response.json()["detail"])

    def test_calendar_age_filter_uses_event_reference_date(self):
        competition = self.competition()
        self.as_role(self.athlete, "athlete")
        response = self.client.get("/api/v1/competitions/?eligible_age=true")
        self.assertEqual([row["id"] for row in response.data], [competition])
        self.athlete.date_of_birth = date(2020, 1, 1)
        self.athlete.save()
        self.assertEqual(self.client.get("/api/v1/competitions/?eligible_age=true").data, [])

    def test_registration_close_reopen_and_download_regulations(self):
        competition = self.competition()
        path = f"/api/v1/competitions/{competition}/registration/"
        self.assertEqual(self.client.post(path, {"registration_open": False}).status_code, 200)
        self.assertEqual(self.apply(competition, self.athlete).status_code, 400)
        self.assertEqual(self.client.post(path, {"registration_open": True}).status_code, 403)
        self.as_role(self.federation, "federation")
        self.assertEqual(self.client.post(path, {"registration_open": True}).status_code, 200)
        self.assertEqual(self.apply(competition, self.athlete).status_code, 201)
        response = self.client.get(f"/api/v1/competitions/{competition}/regulations/")
        self.assertEqual(response.status_code, 200)
        doc = Document(BytesIO(response.content))
        self.assertIn("Тестовый регламент", "\n".join(p.text for p in doc.paragraphs))

    def test_competition_types_manageable_and_role_protected(self):
        self.as_role(self.federation, "federation")
        response = self.client.post("/api/v1/competition-types/", {"name": "Вымышленный кубок"})
        self.assertEqual(response.status_code, 201, response.data)
        pk = response.data["id"]
        self.assertEqual(self.client.patch(f"/api/v1/competition-types/{pk}/", {"is_active": False}, format="json").status_code, 200)
        self.as_role(self.athlete, "athlete")
        self.assertEqual(self.client.post("/api/v1/competition-types/", {"name": "Чужой тип"}).status_code, 403)

    def test_withdraw_renew_history_and_eligibility(self):
        competition = self.competition()
        self.as_role(self.athlete, "athlete")
        self.assertEqual(self.client.post("/api/v1/consents/", {"action": "withdraw"}).status_code, 200)
        self.assertFalse(has_current_consent(self.athlete))
        self.assertEqual(self.apply(competition, self.athlete).status_code, 400)
        self.assertEqual(self.client.post("/api/v1/consents/", {"action": "renew", "consent": True}).status_code, 200)
        self.assertTrue(has_current_consent(self.athlete))
        self.assertEqual(PersonalDataConsent.objects.filter(user=self.athlete).count(), 2)
        self.assertEqual(self.apply(competition, self.athlete).status_code, 201)
        event = AuditEvent.objects.get(action="consent.renewed")
        self.assertEqual(event.active_role, "athlete")
        self.as_role(self.second, "athlete")
        self.assertEqual(len(self.client.get("/api/v1/consents/").data["history"]), 1)

    def test_minor_renewal_requires_guardian_checkbox_and_does_not_verify_identity(self):
        self.athlete.date_of_birth = date(2015, 1, 1)
        self.athlete.save()
        guardian = Guardian.objects.create(minor=self.athlete, full_name="Вымышленный Представитель", relationship="родитель")
        self.as_role(self.athlete, "athlete")
        self.client.post("/api/v1/consents/", {"action": "withdraw"})
        self.assertEqual(self.client.post("/api/v1/consents/", {"action": "renew", "consent": True}).status_code, 400)
        self.assertEqual(self.client.post("/api/v1/consents/", {"action": "renew", "consent": True, "guardian_consent": True}).status_code, 200)
        guardian.refresh_from_db()
        self.assertEqual(guardian.confirmation_status, "pending_verification")

    def test_profile_private_antidoping_validation_and_audit(self):
        self.as_role(self.athlete, "athlete")
        self.assertEqual(self.client.get("/api/v1/profile/").data["organization"], "Без организации")
        self.assertEqual(self.client.patch("/api/v1/profile/", {"antidoping_completed_on": "2026-01-01"}, format="json").status_code, 400)
        data = {"antidoping_completed_on": "2026-01-01", "antidoping_certificate_url": "https://example.test/fictional-certificate"}
        self.assertEqual(self.client.patch("/api/v1/profile/", data, format="json").status_code, 200)
        self.assertTrue(AuditEvent.objects.filter(action="antidoping.updated", active_role="athlete").exists())
        self.as_role(self.coach, "coach")
        self.assertEqual(self.client.get("/api/v1/profile/").status_code, 403)

    def test_link_suspension_revokes_scope_and_resume_requires_athlete(self):
        link = CoachAthleteLink.objects.create(athlete=self.athlete, coach=self.coach, discipline=self.algorithmic, status="confirmed")
        self.as_role(self.coach, "coach")
        path = f"/api/v1/coach-links/{link.pk}/status/"
        self.assertEqual(self.client.post(path, {"status": "suspended", "reason": "Перерыв"}).status_code, 200)
        self.assertEqual(self.client.get(f"/api/v1/reports/athlete-participation/?athlete={self.athlete.pk}").status_code, 403)
        self.assertEqual(self.client.post(path, {"status": "confirmed", "reason": "Возобновление"}).status_code, 400)
        self.as_role(self.second, "athlete")
        self.assertEqual(self.client.post(path, {"status": "ended", "reason": "Чужая связь"}).status_code, 404)
        self.as_role(self.athlete, "athlete")
        self.assertEqual(self.client.post(path, {"status": "confirmed", "reason": "Возобновление"}).status_code, 200)
        self.assertEqual(self.client.post(path, {"status": "ended", "reason": "Завершение"}).status_code, 200)
        self.assertEqual(self.client.post(path, {"status": "confirmed", "reason": "Повтор"}).status_code, 400)
        self.assertEqual(AuditEvent.objects.filter(action="coach_link.status_changed").count(), 3)

    def test_athlete_report_uses_latest_snapshot_private_and_valid_xlsx(self):
        competition = self.competition()
        application = self.apply(competition, self.athlete)
        self.approve(application.data["id"])
        self.client.post(f"/api/v1/competitions/{competition}/protocol-publish/")
        self.client.post(f"/api/v1/competitions/{competition}/protocol-publish/", {"correction_reason": "Вторая версия"})
        self.as_role(self.athlete, "athlete")
        path = "/api/v1/reports/athlete-participation/"
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual(sheet.cell(5, 8).value, 2)
        self.assertEqual(sheet.cell(5, 5).value, "Результат не внесён")
        self.assertIsNone(sheet.cell(5, 6).value)
        self.assertIn("private", response["Cache-Control"])
        self.assertEqual(self.client.get(path + f"?athlete={self.second.pk}").status_code, 403)
        self.as_role(self.coach, "coach")
        self.assertEqual(self.client.get(path + f"?athlete={self.athlete.pk}").status_code, 403)
        CoachAthleteLink.objects.create(athlete=self.athlete, coach=self.coach, discipline=self.algorithmic, status="confirmed")
        self.assertEqual(self.client.get(path + f"?athlete={self.athlete.pk}").status_code, 200)

    def test_empty_report_is_valid_and_formula_name_is_literal(self):
        self.athlete.full_name = "=Вымышленный"
        self.athlete.save()
        self.as_role(self.athlete, "athlete")
        response = self.client.get("/api/v1/reports/athlete-participation/")
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertEqual(sheet.cell(5, 1).value, "Опубликованных результатов пока нет")
