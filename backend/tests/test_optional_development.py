from io import BytesIO

from openpyxl import load_workbook

from apps.competitions.models import CalendarEntry
from apps.extensions.development import safe_cell

from .optional_helpers import OptionalCase


class DevelopmentTests(OptionalCase):
    def test_sources_conducted_definition_and_xlsx(self):
        self.login(self.federation, "federation")
        self.assertFalse(
            any(
                r["kind"] == "Проведённое событие"
                for r in self.client.get("/api/v1/development/").data["items"]
            )
        )
        calendar = CalendarEntry.objects.create(
            source="regional",
            entry_number="DEMO-01",
            title="Демо календарь",
            year=self.competition.starts_at.year,
        )
        self.competition.calendar_entry = calendar
        self.competition.save()
        response = self.client.get("/api/v1/development/")
        event = next(
            r for r in response.data["items"] if r["kind"] == "Проведённое событие"
        )
        self.assertIn(f"ProtocolPublication:{self.publication.pk}", event["source"])
        self.assertFalse(response.data["registry_available"])
        xlsx = self.client.get("/api/v1/development/?file_format=xlsx")
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn("Методика", load_workbook(BytesIO(xlsx.content)).sheetnames)
        self.assertEqual(safe_cell("=HYPERLINK(1)"), "'=HYPERLINK(1)")
        self.assertEqual(
            self.client.get(
                "/api/v1/development/?start_year=2026&end_year=2026"
            ).status_code,
            400,
        )
        self.login(self.athlete, "athlete")
        self.assertEqual(self.client.get("/api/v1/development/").status_code, 403)
