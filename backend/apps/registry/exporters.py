from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def safe_text(value):
    value = str(value)
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def export_participation(rows, filters, fingerprint):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Итоги"
    summary.append(["Численность занимающихся спортивным программированием в Оренбургской области"])
    summary.append(["Годы", f"{rows[0]['year']}–{rows[-1]['year']}", "Время формирования", timezone.localtime().isoformat(timespec="seconds")])
    summary.append(["Год", "Уникальных спортсменов", "Начали", "Завершили", "Изменение", "Рост, %"])
    for row in rows:
        summary.append([row[key] for key in ("year", "unique_athletes", "started", "ceased", "change", "growth_percent")])
    for key, title in (("organizations", "Организации"), ("disciplines", "Дисциплины"), ("municipalities", "Муниципалитеты"), ("age_groups", "Возрастные группы")):
        sheet = workbook.create_sheet(title)
        sheet.append(["Год", title, "Участников; категории не суммируются"])
        for row in rows:
            for item in row["breakdowns"][key]:
                sheet.append([row["year"], safe_text(item["name"]), item["participants"]])
    methodology = workbook.create_sheet("Методика")
    methodology.append(["Отпечаток содержимого отчёта SHA-256: " + fingerprint])
    methodology.append([safe_text("Фильтры: " + str(filters))])
    methodology.append(["Уникальные спортсмены: отдельные люди с периодом участия, пересекающим календарный год и все выбранные фильтры одновременно."])
    methodology.append(["Исторические организации, муниципалитеты и дисциплины учитываются по пересечению периодов; суммы категорий не равны числу людей."])
    methodology.append(["Возраст на 31 декабря года; границы групп из справочника. Начало и окончание — границы периодов в году. Данные не являются официальной статистикой."])
    for sheet in workbook:
        sheet.freeze_panes = "A4" if sheet == summary else "A2"
        sheet.column_dimensions["A"].width = 18 if sheet != methodology else 100
        sheet.column_dimensions["B"].width = 50
        for col in "CDEF":
            sheet.column_dimensions[col].width = 24
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="221F73")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
