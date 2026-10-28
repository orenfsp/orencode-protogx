from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.competitions.exporters import _font_path
from apps.qualifications.models import SportRankHistory


def _rows(team):
    return list(team.memberships.filter(removed_at__isnull=True).select_related("athlete", "athlete__athlete_profile__organization").order_by("position"))


def _organization(athlete):
    profile = getattr(athlete, "athlete_profile", None)
    return profile.organization.name if profile and profile.organization else "Не указана"


def _title(team, draft=False):
    return ("ПРОЕКТ ПРИКАЗА" if draft else "СОСТАВ РЕГИОНАЛЬНОЙ КОМАНДЫ")


def _rank(team, athlete):
    record = SportRankHistory.objects.filter(athlete=athlete, discipline=team.discipline, assigned_on__lte=timezone.localdate(), valid_until__gte=timezone.localdate()).order_by("-assigned_on", "-pk").first()
    return record.get_rank_display() if record else "—"


def render_team_docx(team, *, draft=False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "DejaVu Sans"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(_title(team, draft=draft), level=0)
    doc.add_paragraph(f"Команда: {team.name}")
    doc.add_paragraph(f"Регион: {team.region}")
    doc.add_paragraph(f"Сезон: {team.season.year}; дисциплина: {team.discipline.name}; возрастная группа: {team.age_group.name}")
    doc.add_paragraph(f"Сформировано: {timezone.localtime():%d.%m.%Y %H:%M %Z}")
    if draft:
        doc.add_paragraph("Проект: предлагается утвердить приведённый ниже состав. Документ не является подписанным официальным приказом; форма требует проверки.")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Cm(0.9), Cm(5.2), Cm(5.3), Cm(3.2), Cm(2.1))
    for column, width in zip(table.columns, widths):
        column.width = width
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for cell, text in zip(table.rows[0].cells, ("№", "Спортсмен", "Организация", "Спортивный разряд", "Рейтинг")):
        cell.text = text
    for index, membership in enumerate(_rows(team), 1):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for cell, width in zip(cells, widths):
            cell.width = width
        cells[0].text = str(index)
        cells[1].text = membership.athlete.full_name
        cells[2].text = _organization(membership.athlete)
        cells[3].text = _rank(team, membership.athlete)
        cells[4].text = membership.eligibility_snapshot.get("rating_score") or "—"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph("Подпись представителя федерации: ____________________")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def render_team_pdf(team, *, draft=False):
    font_name = "FSP-DejaVu"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, _font_path()))
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=_title(team, draft=draft), invariant=1)
    normal = ParagraphStyle("team-normal", fontName=font_name, fontSize=9, leading=13)
    heading = ParagraphStyle("team-heading", parent=normal, fontSize=16, leading=20, textColor=colors.HexColor("#221F73"), spaceAfter=8)
    story = [Paragraph(_title(team, draft=draft), heading), Paragraph(escape(f"Команда: {team.name}"), normal), Paragraph(escape(f"Регион: {team.region}; сезон: {team.season.year}; дисциплина: {team.discipline.name}; возрастная группа: {team.age_group.name}"), normal), Paragraph(f"Сформировано: {timezone.localtime():%d.%m.%Y %H:%M %Z}", normal), Spacer(1, 10)]
    if draft:
        story.append(Paragraph("Предлагается утвердить приведённый ниже состав. Форма требует проверки.", normal))
    data = [["№", "Спортсмен", "Организация", "Разряд", "Рейтинг"]]
    data.extend([[str(index), Paragraph(escape(item.athlete.full_name), normal), Paragraph(escape(_organization(item.athlete)), normal), Paragraph(escape(_rank(team, item.athlete)), normal), item.eligibility_snapshot.get("rating_score") or "—"] for index, item in enumerate(_rows(team), 1)])
    table = Table(data, colWidths=[9 * mm, 49 * mm, 50 * mm, 34 * mm, 20 * mm], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEDF9")), ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#221F73")), ("FONTNAME", (0, 0), (-1, -1), font_name), ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.extend([table, Spacer(1, 12), Paragraph("Проект не является подписанным официальным приказом." if draft else "Сформировано из базы данных платформы.", normal), Paragraph("Подпись представителя федерации: ____________________", normal)])
    doc.build(story)
    return stream.getvalue()
