from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, extend_schema
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.views import APIView

from apps.accounts.models import User
from apps.registry.exporters import safe_text

from .exporters import render_protocol_docx, render_protocol_pdf
from .models import CoachAthleteLink, ProtocolPublication
from .services import require_role

STATUS = {"classified": "Классифицирован", "no_show": "Не явился", "unclassified": "Без результата", "disqualified": "Дисквалифицирован", "pending_result": "Результат не внесён"}


def athlete_report_rows(athlete, discipline_ids=None):
    publications = ProtocolPublication.objects.filter(competition__applications__athlete=athlete, competition__applications__status="approved").select_related("competition__discipline").distinct().order_by("competition_id", "-version")
    seen, rows = set(), []
    for publication in publications:
        if publication.competition_id in seen:
            continue
        seen.add(publication.competition_id)
        if discipline_ids is not None and publication.competition.discipline_id not in discipline_ids:
            continue
        for row in publication.snapshot["rows"]:
            if row["athlete_id"] == athlete.pk:
                rows.append((publication, row))
    return rows


class AthleteReportView(APIView):
    @extend_schema(parameters=[OpenApiParameter("athlete", int)], responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY})
    def get(self, request):
        role = require_role(request, "athlete", "coach", "federation")
        try:
            athlete_id = int(request.query_params.get("athlete", request.user.pk))
        except (TypeError, ValueError) as error:
            raise ValidationError("Некорректный идентификатор спортсмена") from error
        if role == "athlete" and athlete_id != request.user.pk:
            raise PermissionDenied("Доступен только собственный отчёт")
        disciplines = None
        if role == "coach":
            disciplines = set(CoachAthleteLink.objects.filter(coach=request.user, athlete_id=athlete_id, status="confirmed").values_list("discipline_id", flat=True))
            if not disciplines:
                raise PermissionDenied("Нужна подтверждённая связь со спортсменом")
        athlete = get_object_or_404(User, pk=athlete_id, role_assignments__role="athlete")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Выступления"
        sheet.append(["Отчёт о соревнованиях спортсмена — демонстрационная форма"])
        sheet.append(["Спортсмен", safe_text(athlete.full_name)])
        sheet.append(["Сформирован", timezone.localtime().isoformat(), "Источник: последние опубликованные версии протоколов"])
        sheet.append(["Соревнование", "Дисциплина", "Дата", "Уровень", "Статус результата", "Место", "Баллы", "Версия протокола", "Публикация"])
        for publication, row in athlete_report_rows(athlete, disciplines):
            metadata = publication.snapshot.get("metadata", {})
            sheet.append([safe_text(publication.snapshot["competition_title"]), safe_text(publication.snapshot["discipline"]), metadata.get("starts_at", ""), safe_text(metadata.get("level", "")), STATUS.get(row["status"], "Не определён"), row.get("place"), safe_text(row["score"]) if row.get("score") is not None else None, publication.version, publication.pk])
        if sheet.max_row == 4:
            sheet.append(["Опубликованных результатов пока нет"])
        sheet.merge_cells("A1:I1")
        sheet.row_dimensions[1].height = 35
        sheet.merge_cells("B2:I2")
        sheet.merge_cells("C3:I3")
        sheet.auto_filter.ref = f"A4:I{sheet.max_row}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:4"
        sheet.freeze_panes = "A5"
        for column in "ABCDEFGHI":
            sheet.column_dimensions[column].width = 30 if column not in "FGHI" else 18
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[4]:
            cell.font = Font(bold=True)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="athlete-{athlete.pk}-participation.xlsx"'
        response["Cache-Control"] = "private, no-store"
        return response


class PersonalProtocolView(APIView):
    @extend_schema(parameters=[OpenApiParameter("file_format", str, enum=["pdf", "docx"])], responses={(200, "application/pdf"): OpenApiTypes.BINARY, (200, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"): OpenApiTypes.BINARY})
    def get(self, request, pk, version):
        require_role(request, "athlete")
        publication = get_object_or_404(ProtocolPublication, competition_id=pk, version=version)
        snapshot = deepcopy(publication.snapshot)
        snapshot["rows"] = [row for row in snapshot["rows"] if row["athlete_id"] == request.user.pk]
        if not snapshot["rows"]:
            raise PermissionDenied("В этой версии нет ваших результатов")
        snapshot["competition_title"] = "Личная выписка · " + snapshot["competition_title"]
        snapshot["sections"] = [item for item in snapshot["sections"] if item not in ("judges", "signatures")]
        snapshot["judges"] = []
        snapshot["signature_roles"] = []
        excerpt = SimpleNamespace(snapshot=snapshot, version=publication.version, published_at=publication.published_at, correction_reason=publication.correction_reason)
        output_format = request.query_params.get("file_format", "pdf")
        if output_format not in ("pdf", "docx"):
            raise ValidationError("Формат: pdf или docx")
        response = HttpResponse(render_protocol_pdf(excerpt) if output_format == "pdf" else render_protocol_docx(excerpt), content_type="application/pdf" if output_format == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        response["Content-Disposition"] = f'attachment; filename="personal-protocol-{pk}-v{version}.{output_format}"'
        response["Cache-Control"] = "private, no-store"
        return response
