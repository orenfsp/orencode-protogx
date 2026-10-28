from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from docx import Document
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import extend_schema
from rest_framework import serializers
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import Competition
from .serializers import CompetitionSerializer
from .services import audit, require_role
from .views import CompetitionDetailView


class RegistrationInputSerializer(serializers.Serializer):
    registration_open = serializers.BooleanField()


class RegistrationView(APIView):
    @extend_schema(request=RegistrationInputSerializer, responses=CompetitionSerializer)
    @transaction.atomic
    def post(self, request, pk):
        require_role(request, "federation")
        competition = get_object_or_404(Competition.objects.select_for_update(), pk=pk)
        serializer = RegistrationInputSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        value = serializer.validated_data["registration_open"]
        if competition.state != "published" or (value and timezone.now() >= competition.registration_deadline):
            raise ValidationError("Приём заявок доступен только до срока в опубликованном соревновании")
        before = competition.registration_open
        competition.registration_open = value
        competition.save(update_fields=["registration_open", "updated_at"])
        audit(request.user, "competition.registration_changed", competition, before={"registration_open": before}, after={"registration_open": value})
        return Response(CompetitionSerializer(competition).data)


class RegulationsDownloadView(APIView):
    permission_classes = CompetitionDetailView.permission_classes

    @extend_schema(responses={(200, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"): OpenApiTypes.BINARY})
    def get(self, request, pk):
        competition = CompetitionDetailView().get_object(request, pk)
        document = Document()
        document.add_heading("Положение (регламент)", 0)
        document.add_heading(competition.title, 1)
        document.add_paragraph(competition.discipline.name)
        for paragraph in competition.regulations.splitlines():
            document.add_paragraph(paragraph)
        document.add_paragraph(f"Редакция в платформе: {competition.updated_at.isoformat()}")
        output = BytesIO()
        document.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        response["Content-Disposition"] = f'attachment; filename="regulations-{pk}.docx"'
        response["Cache-Control"] = "private, no-store"
        return response
