"""Render an immutable published snapshot as clearly unverified protocol documents."""

from io import BytesIO
from pathlib import Path
from textwrap import wrap
from xml.sax.saxutils import escape

from django.conf import settings
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INDIGO = "#221F73"
ORANGE = "#FF9900"
STATUS = {
    "classified": "Классифицирован",
    "no_show": "Не явился",
    "unclassified": "Без результата",
    "disqualified": "Дисквалифицирован",
    "pending_result": "Результат не внесён",
}


def _font_path():
    candidates = [
        Path(getattr(settings, "PROTOCOL_PDF_FONT", "")),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    raise RuntimeError("A Cyrillic DejaVu Sans font is required for protocol PDF export")


def _cell(row, key):
    value = row.get(key)
    if key == "status":
        return STATUS.get(value, str(value))
    return "—" if value is None or value == "" else str(value)


def _heading(publication):
    snapshot = publication.snapshot
    return [
        "Демонстрационный протокол — требуется проверка первоисточника",
        snapshot["competition_title"],
        f"Версия {publication.version} · опубликовано {publication.published_at:%d.%m.%Y %H:%M}",
    ]


def _metadata_lines(snapshot):
    meta = snapshot["metadata"]
    return [
        f"Дисциплина: {snapshot['discipline']}",
        f"Вид и уровень: {meta['competition_type']} · {meta['level']}",
        f"Дата и место: {meta['starts_at'][:10]} · {meta['location']}",
        f"Формат: {meta['format']} · возраст на дату {meta['age_reference_date']}",
        f"Календарная запись: {meta['calendar_entry'] or 'не указана'}",
    ]


def _register_font():
    if "FSP-DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("FSP-DejaVu", _font_path()))


def render_protocol_pdf(publication):
    _register_font()
    snapshot = publication.snapshot
    stream = BytesIO()
    page_size = landscape(A3 if len(snapshot["columns"]) > 8 else A4)
    document = SimpleDocTemplate(stream, pagesize=page_size, leftMargin=38, rightMargin=38, topMargin=38, bottomMargin=42, title=f"Протокол {snapshot['competition_title']} v{publication.version}", author="ФСП Оренбургская область", invariant=1)
    normal = ParagraphStyle("normal", fontName="FSP-DejaVu", fontSize=8.5, leading=12, wordWrap="CJK", textColor=colors.HexColor("#1B1C21"))
    title = ParagraphStyle("title", parent=normal, fontSize=17, leading=22, textColor=colors.HexColor(INDIGO), spaceAfter=8)
    warning = ParagraphStyle("warning", parent=normal, fontSize=9, leading=14, textColor=colors.HexColor("#1B1C21"), backColor=colors.HexColor("#FFF1D6"), borderPadding=7, spaceAfter=14)
    small = ParagraphStyle("small", parent=normal, fontSize=7.5, leading=10)
    story = [Paragraph(escape(_heading(publication)[0]), warning), Paragraph(escape(_heading(publication)[1]), title), Paragraph(escape(_heading(publication)[2]), normal), Spacer(1, 12)]
    story.extend(Paragraph(escape(line), normal) for line in _metadata_lines(snapshot))
    if publication.correction_reason:
        story.append(Paragraph(escape(f"Причина изменения: {publication.correction_reason}"), normal))
    story.append(Spacer(1, 14))
    columns = snapshot["columns"]
    data = [[Paragraph(escape(column["label"]), small) for column in columns]]
    for row in snapshot["rows"]:
        data.append([Paragraph(escape(_cell(row, column["key"])), small) for column in columns])
    width = page_size[0] - 76
    table = Table(data, colWidths=[width / len(columns)] * len(columns), repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEDF9")),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor(INDIGO)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([table, Spacer(1, 16)])
    if "judges" in snapshot["sections"]:
        story.append(Paragraph("Судейская коллегия", title))
        story.extend(Paragraph(escape(f"{judge['role']}: {judge['name']}"), normal) for judge in snapshot["judges"])
    if "signatures" in snapshot["sections"]:
        story.extend([Spacer(1, 10), Paragraph("Поля подписей — не электронная подпись", normal)])
        story.extend(Paragraph(escape(f"{role}: ____________________"), normal) for role in snapshot["signature_roles"])
    story.extend([Spacer(1, 12), Paragraph(escape(snapshot["source_note"] or "Форма требует проверки по первоисточнику."), small)])
    if snapshot["source_url"]:
        for line in wrap(snapshot["source_url"], width=100, break_long_words=True, break_on_hyphens=False):
            story.append(Paragraph(escape(line), small))

    def footer(canvas, _doc):
        canvas.setFont("FSP-DejaVu", 7)
        canvas.setFillColor(colors.HexColor(INDIGO))
        canvas.drawString(38, 21, f"Версия {publication.version} · демонстрационный документ")
        canvas.drawRightString(page_size[0] - 38, 21, f"Страница {_doc.page}")

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def render_protocol_docx(publication):
    snapshot = publication.snapshot
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(9)
    warning = doc.add_paragraph(_heading(publication)[0])
    warning.style = doc.styles["Subtitle"]
    warning.runs[0].font.color.rgb = RGBColor(0x22, 0x1F, 0x73)
    doc.add_heading(_heading(publication)[1], level=0)
    doc.add_paragraph(_heading(publication)[2])
    for line in _metadata_lines(snapshot):
        doc.add_paragraph(line)
    if publication.correction_reason:
        doc.add_paragraph(f"Причина изменения: {publication.correction_reason}")
    columns = snapshot["columns"]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = column["label"]
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0x22, 0x1F, 0x73)
    for row in snapshot["rows"]:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cells[index].text = _cell(row, column["key"])
    if "judges" in snapshot["sections"]:
        doc.add_heading("Судейская коллегия", level=2)
        for judge in snapshot["judges"]:
            doc.add_paragraph(f"{judge['role']}: {judge['name']}")
    if "signatures" in snapshot["sections"]:
        doc.add_paragraph("Поля подписей — не электронная подпись")
        for role in snapshot["signature_roles"]:
            doc.add_paragraph(f"{role}: ____________________")
    doc.add_paragraph(snapshot["source_note"] or "Форма требует проверки по первоисточнику.")
    if snapshot["source_url"]:
        doc.add_paragraph(snapshot["source_url"])
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
